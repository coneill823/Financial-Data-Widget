"""
Financial Data Dashboard
Loads labor and expense data from Excel and visualises spend vs budget.
"""

import os
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure

# ── palette ──────────────────────────────────────────────────────────────────
C = {
    "bg":       "#1e1e2e",
    "panel":    "#2a2a3e",
    "border":   "#3a3a5e",
    "text":     "#e2e8f0",
    "dim":      "#94a3b8",
    "accent":   "#7c3aed",
    "cyan":     "#06b6d4",
    "green":    "#10b981",
    "yellow":   "#f59e0b",
    "red":      "#ef4444",
    "cat": {
        "Material Costs":     "#7c3aed",
        "Purchased Services": "#06b6d4",
        "T&L":                "#10b981",
        "Labor":              "#f59e0b",
    },
}
FALLBACK_COLORS = ["#a78bfa", "#67e8f9", "#6ee7b7", "#fcd34d",
                   "#f87171", "#c084fc", "#34d399", "#fb923c"]


# ── helpers ───────────────────────────────────────────────────────────────────

def _fmt(value: float) -> str:
    return f"${value:,.2f}"


def _style_ax(ax):
    ax.set_facecolor(C["panel"])
    ax.tick_params(colors=C["text"], labelsize=8)
    ax.xaxis.label.set_color(C["text"])
    ax.yaxis.label.set_color(C["text"])
    ax.title.set_color(C["text"])
    for sp in ax.spines.values():
        sp.set_color(C["border"])


def _money_formatter(x, _):
    if abs(x) >= 1_000_000:
        return f"${x/1_000_000:.1f}M"
    if abs(x) >= 1_000:
        return f"${x/1_000:.0f}k"
    return f"${x:.0f}"


# ── main app ──────────────────────────────────────────────────────────────────

class FinancialDashboard:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("Financial Data Dashboard")
        self.root.geometry("1440x920")
        self.root.configure(bg=C["bg"])

        self.labor_df:    pd.DataFrame | None = None
        self.expenses_df: pd.DataFrame | None = None
        self.budgets_df:  pd.DataFrame | None = None
        self.projects:    list[str] = []

        self._canvas_widget = None
        self._fig: Figure | None = None

        self._build_ui()

    # ── UI construction ───────────────────────────────────────────────────────

    def _build_ui(self):
        # ── header ────────────────────────────────────────────────────────────
        hdr = tk.Frame(self.root, bg=C["bg"])
        hdr.pack(fill="x", padx=20, pady=(14, 6))

        tk.Label(hdr, text="Financial Data Dashboard",
                 font=("Arial", 22, "bold"),
                 bg=C["bg"], fg=C["text"]).pack(side="left")

        # ── toolbar ───────────────────────────────────────────────────────────
        tb = tk.Frame(self.root, bg=C["panel"], pady=10, padx=16)
        tb.pack(fill="x", padx=20, pady=(0, 10))

        self._btn(tb, "Load Excel File", self._load_excel,
                  bg=C["accent"]).pack(side="left", padx=(0, 10))
        self._btn(tb, "Download Template", self._create_template,
                  bg=C["border"]).pack(side="left", padx=(0, 24))

        tk.Label(tb, text="View:", bg=C["panel"], fg=C["text"],
                 font=("Arial", 11)).pack(side="left", padx=(0, 6))

        self.project_var = tk.StringVar(value="All Projects")
        self.project_combo = ttk.Combobox(
            tb, textvariable=self.project_var,
            values=["All Projects"], state="readonly",
            width=28, font=("Arial", 11))
        self.project_combo.pack(side="left", padx=(0, 24))
        self.project_combo.bind("<<ComboboxSelected>>", self._refresh)

        self.file_lbl = tk.Label(tb, text="No file loaded",
                                 bg=C["panel"], fg=C["dim"],
                                 font=("Arial", 10))
        self.file_lbl.pack(side="left")

        # ── summary bar ───────────────────────────────────────────────────────
        self.summary_bar = tk.Frame(self.root, bg=C["panel"])
        self.summary_bar.pack(fill="x", padx=20, pady=(0, 8))

        # ── chart canvas area ─────────────────────────────────────────────────
        self.chart_area = tk.Frame(self.root, bg=C["bg"])
        self.chart_area.pack(fill="both", expand=True, padx=20, pady=(0, 16))

        self._show_placeholder()

    def _btn(self, parent, text, cmd, bg=None):
        return tk.Button(parent, text=text, command=cmd,
                         bg=bg or C["accent"], fg="white",
                         font=("Arial", 11, "bold"),
                         relief="flat", padx=14, pady=5,
                         activebackground=C["border"],
                         activeforeground=C["text"],
                         cursor="hand2")

    def _show_placeholder(self):
        for w in self.chart_area.winfo_children():
            w.destroy()
        tk.Label(self.chart_area,
                 text=("Load an Excel file to view financial data\n\n"
                       "Click  'Download Template'  to get the expected format"),
                 bg=C["bg"], fg=C["dim"],
                 font=("Arial", 14), justify="center").pack(expand=True)

    # ── Excel loading ─────────────────────────────────────────────────────────

    def _load_excel(self):
        path = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")])
        if not path:
            return

        try:
            xl = pd.ExcelFile(path)
            names = xl.sheet_names

            def find(keywords):
                for s in names:
                    if any(k in s.lower() for k in keywords):
                        return s
                return None

            labor_sheet   = find(["labor", "hours", "people", "timesheet"])
            expense_sheet = find(["expense", "cost", "spend"])
            budget_sheet  = find(["budget", "project"])

            if not labor_sheet or not expense_sheet:
                messagebox.showerror(
                    "Sheet Not Found",
                    f"Expected sheets named 'Labor' and 'Expenses'.\n"
                    f"Found: {', '.join(names)}")
                return

            self.labor_df    = self._parse_labor(xl, labor_sheet)
            self.expenses_df = self._parse_expenses(xl, expense_sheet)
            self.budgets_df  = self._parse_budgets(xl, budget_sheet) if budget_sheet else None

            projects = set()
            for df in (self.labor_df, self.expenses_df):
                if df is not None and "project" in df.columns:
                    projects.update(df["project"].dropna().unique())

            self.projects = sorted(projects)
            self.project_combo["values"] = ["All Projects"] + self.projects
            self.project_var.set("All Projects")

            self.file_lbl.config(
                text=f"Loaded: {os.path.basename(path)}", fg=C["green"])
            self._refresh()

        except Exception as exc:
            messagebox.showerror("Load Error", str(exc))
            import traceback; traceback.print_exc()

    # -- parsers ---------------------------------------------------------------

    @staticmethod
    def _normalise_cols(df: pd.DataFrame) -> pd.DataFrame:
        df.columns = [c.strip().lower().replace(" ", "_") for c in df.columns]
        return df

    def _parse_labor(self, xl, sheet) -> pd.DataFrame:
        df = self._normalise_cols(pd.read_excel(xl, sheet_name=sheet))

        date_col = next((c for c in df.columns if "date" in c), None)
        if date_col:
            df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
            df.rename(columns={date_col: "date"}, inplace=True)

        proj_col = next((c for c in df.columns if "project" in c), None)
        if proj_col and proj_col != "project":
            df.rename(columns={proj_col: "project"}, inplace=True)

        hour_col = next((c for c in df.columns if "hour" in c), None)
        rate_col = next((c for c in df.columns if any(
            x in c for x in ["rate", "hourly", "salary", "wage"])), None)

        if hour_col:
            df.rename(columns={hour_col: "hours"}, inplace=True)
        if rate_col:
            df.rename(columns={rate_col: "rate"}, inplace=True)

        if "hours" in df.columns:
            rate = df["rate"] if "rate" in df.columns else 50.0
            df["labor_cost"] = df["hours"] * rate

        return df

    def _parse_expenses(self, xl, sheet) -> pd.DataFrame:
        df = self._normalise_cols(pd.read_excel(xl, sheet_name=sheet))

        date_col = next((c for c in df.columns if "date" in c), None)
        if date_col:
            df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
            df.rename(columns={date_col: "date"}, inplace=True)

        for std, kws in [
            ("project",  ["project"]),
            ("cost",     ["cost", "amount", "value", "spend", "price"]),
            ("category", ["categor", "type"]),
        ]:
            col = next((c for c in df.columns if any(k in c for k in kws)), None)
            if col and col != std:
                df.rename(columns={col: std}, inplace=True)

        if "cost" in df.columns:
            df["cost"] = pd.to_numeric(df["cost"], errors="coerce").fillna(0)

        return df

    def _parse_budgets(self, xl, sheet) -> pd.DataFrame | None:
        try:
            df = self._normalise_cols(pd.read_excel(xl, sheet_name=sheet))
            proj_col   = next((c for c in df.columns if "project" in c), None)
            budget_col = next((c for c in df.columns
                               if any(x in c for x in ["budget", "max", "limit"])), None)
            if proj_col and budget_col:
                df.rename(columns={proj_col: "project", budget_col: "budget"}, inplace=True)
                return df[["project", "budget"]]
        except Exception:
            pass
        return None

    # ── data helpers ──────────────────────────────────────────────────────────

    def _filtered(self) -> tuple[pd.DataFrame, pd.DataFrame]:
        sel = self.project_var.get()
        labor    = self.labor_df.copy()    if self.labor_df    is not None else pd.DataFrame()
        expenses = self.expenses_df.copy() if self.expenses_df is not None else pd.DataFrame()

        if sel != "All Projects":
            if not labor.empty and "project" in labor.columns:
                labor = labor[labor["project"] == sel]
            if not expenses.empty and "project" in expenses.columns:
                expenses = expenses[expenses["project"] == sel]
        return labor, expenses

    def _get_budget(self, project: str) -> float | None:
        if self.budgets_df is None or "budget" not in self.budgets_df.columns:
            return None
        if project == "All Projects":
            return float(self.budgets_df["budget"].sum())
        row = self.budgets_df[self.budgets_df["project"] == project]
        return float(row["budget"].iloc[0]) if not row.empty else None

    def _project_totals(self) -> pd.DataFrame:
        rows = []
        for proj in self.projects:
            lc = (self.labor_df[self.labor_df["project"] == proj]["labor_cost"].sum()
                  if self.labor_df is not None and "project" in self.labor_df.columns
                     and "labor_cost" in self.labor_df.columns else 0)
            ec = (self.expenses_df[self.expenses_df["project"] == proj]["cost"].sum()
                  if self.expenses_df is not None and "project" in self.expenses_df.columns
                     and "cost" in self.expenses_df.columns else 0)
            bud = self._get_budget(proj) or 0
            rows.append({"project": proj, "labor": lc, "expenses": ec,
                         "total": lc + ec, "budget": bud})
        return pd.DataFrame(rows)

    # ── refresh / drawing ─────────────────────────────────────────────────────

    def _refresh(self, _event=None):
        if self.labor_df is None:
            return
        labor, expenses = self._filtered()
        sel = self.project_var.get()

        self._update_summary(labor, expenses, sel)
        self._draw_charts(labor, expenses, sel)

    def _update_summary(self, labor, expenses, sel):
        for w in self.summary_bar.winfo_children():
            w.destroy()

        labor_total   = labor["labor_cost"].sum()   if "labor_cost" in labor.columns   else 0
        expense_total = expenses["cost"].sum()       if "cost" in expenses.columns       else 0
        total         = labor_total + expense_total
        budget        = self._get_budget(sel)

        cards = [
            ("Total Spend",    _fmt(total),         C["text"]),
            ("Labor Cost",     _fmt(labor_total),   C["cyan"]),
            ("Expenses",       _fmt(expense_total), C["accent"]),
        ]
        if budget:
            remaining = budget - total
            pct = (total / budget * 100) if budget else 0
            cards += [
                ("Max Budget",  _fmt(budget),     C["yellow"]),
                ("Remaining",   _fmt(remaining),  C["green"] if remaining >= 0 else C["red"]),
                ("% Used",      f"{pct:.1f}%",    C["yellow"] if pct < 90 else C["red"]),
            ]

        for label, value, color in cards:
            card = tk.Frame(self.summary_bar, bg=C["bg"], padx=18, pady=6)
            card.pack(side="left", padx=8)
            tk.Label(card, text=label, bg=C["bg"], fg=C["dim"],
                     font=("Arial", 9)).pack()
            tk.Label(card, text=value, bg=C["bg"], fg=color,
                     font=("Arial", 15, "bold")).pack()

    def _draw_charts(self, labor, expenses, sel):
        # destroy old canvas cleanly
        for w in self.chart_area.winfo_children():
            w.destroy()
        if self._fig is not None:
            plt.close(self._fig)

        self._fig = Figure(figsize=(14, 7.5), facecolor=C["bg"])
        ax1 = self._fig.add_subplot(2, 2, 1)
        ax2 = self._fig.add_subplot(2, 2, 2)
        ax3 = self._fig.add_subplot(2, 2, 3)
        ax4 = self._fig.add_subplot(2, 2, 4)

        for ax in (ax1, ax2, ax3, ax4):
            _style_ax(ax)

        self._chart_cumulative(ax1, labor, expenses)
        self._chart_categories(ax2, labor, expenses)
        self._chart_budget_util(ax3, sel)
        self._chart_monthly(ax4, labor, expenses)

        self._fig.tight_layout(pad=2.5)

        canvas = FigureCanvasTkAgg(self._fig, master=self.chart_area)
        canvas.draw()
        canvas.get_tk_widget().pack(fill="both", expand=True)

    # ── chart 1 : cumulative spend ────────────────────────────────────────────

    def _chart_cumulative(self, ax, labor, expenses):
        ax.set_title("Cumulative Spend Over Time", pad=8)
        ax.yaxis.set_major_formatter(plt.FuncFormatter(_money_formatter))

        series_defs: list[tuple[str, pd.Series, str]] = []

        if "labor_cost" in labor.columns and "date" in labor.columns:
            ts = labor.groupby("date")["labor_cost"].sum().sort_index().cumsum()
            series_defs.append(("Labor", ts, C["cat"]["Labor"]))

        if "cost" in expenses.columns and "date" in expenses.columns:
            ts = expenses.groupby("date")["cost"].sum().sort_index().cumsum()
            series_defs.append(("Expenses", ts, C["accent"]))

        if not series_defs:
            ax.text(0.5, 0.5, "No date data available",
                    transform=ax.transAxes, ha="center", va="center",
                    color=C["dim"], fontsize=10)
            return

        min_date = min(s.index.min() for _, s, _ in series_defs)
        max_date = max(s.index.max() for _, s, _ in series_defs)
        idx = pd.date_range(min_date, max_date, freq="D")
        total = pd.Series(0.0, index=idx)

        for name, s, color in series_defs:
            ax.plot(s.index, s.values, label=name, color=color,
                    linewidth=1.8, alpha=0.9)
            total += s.reindex(idx).ffill().fillna(0)

        ax.plot(idx, total.values, label="Total", color="white",
                linewidth=2.2, linestyle="--")
        ax.fill_between(idx, total.values, alpha=0.08, color="white")

        ax.legend(fontsize=7, facecolor=C["bg"], labelcolor=C["text"],
                  edgecolor=C["border"])
        ax.set_xlabel("Date", fontsize=8)
        ax.set_ylabel("Cumulative Spend", fontsize=8)
        plt.setp(ax.xaxis.get_majorticklabels(), rotation=30, ha="right", fontsize=7)

    # ── chart 2 : category pie ────────────────────────────────────────────────

    def _chart_categories(self, ax, labor, expenses):
        ax.set_title("Spend by Category", pad=8)

        data: dict[str, float] = {}

        if "category" in expenses.columns and "cost" in expenses.columns:
            for cat, val in expenses.groupby("category")["cost"].sum().items():
                data[str(cat)] = float(val)

        if "labor_cost" in labor.columns:
            data["Labor"] = float(labor["labor_cost"].sum())

        data = {k: v for k, v in data.items() if v > 0}

        if not data:
            ax.text(0.5, 0.5, "No category data",
                    transform=ax.transAxes, ha="center", va="center",
                    color=C["dim"], fontsize=10)
            return

        labels = list(data.keys())
        values = list(data.values())
        colors = [C["cat"].get(l, FALLBACK_COLORS[i % len(FALLBACK_COLORS)])
                  for i, l in enumerate(labels)]

        wedges, _, autotexts = ax.pie(
            values, colors=colors, autopct="%1.1f%%",
            startangle=90, pctdistance=0.78,
            wedgeprops={"linewidth": 1.5, "edgecolor": C["bg"]})

        for at in autotexts:
            at.set_color("white")
            at.set_fontsize(8)

        ax.legend(wedges,
                  [f"{l}  {_fmt(v)}" for l, v in zip(labels, values)],
                  loc="center left", bbox_to_anchor=(-0.35, 0.5),
                  fontsize=7, facecolor=C["bg"],
                  labelcolor=C["text"], edgecolor=C["border"])

    # ── chart 3 : budget utilisation ──────────────────────────────────────────

    def _chart_budget_util(self, ax, sel):
        ax.set_title("Spend vs Max Budget by Project", pad=8)
        ax.yaxis.set_major_formatter(plt.FuncFormatter(_money_formatter))

        df = self._project_totals()
        if df.empty:
            ax.text(0.5, 0.5, "No project data",
                    transform=ax.transAxes, ha="center", va="center",
                    color=C["dim"], fontsize=10)
            return

        # highlight selected project
        if sel != "All Projects":
            df = df[df["project"] == sel] if sel in df["project"].values else df

        df = df.sort_values("total", ascending=False).head(12)
        x = np.arange(len(df))
        w = 0.38

        bars_spent = ax.bar(x - w / 2, df["total"], w,
                            label="Total Spent", color=C["accent"], alpha=0.9)
        if df["budget"].any():
            ax.bar(x + w / 2, df["budget"], w,
                   label="Max Budget", color=C["yellow"], alpha=0.5)

        ax.set_xticks(x)
        ax.set_xticklabels(
            [p[:14] + "…" if len(p) > 14 else p for p in df["project"]],
            rotation=30, ha="right", fontsize=7)
        ax.set_ylabel("Amount", fontsize=8)
        ax.legend(fontsize=7, facecolor=C["bg"], labelcolor=C["text"],
                  edgecolor=C["border"])

        for bar in bars_spent:
            h = bar.get_height()
            if h > 0:
                ax.annotate(_fmt(h),
                            xy=(bar.get_x() + bar.get_width() / 2, h),
                            xytext=(0, 3), textcoords="offset points",
                            ha="center", va="bottom",
                            fontsize=6, color=C["text"])

    # ── chart 4 : monthly stacked bar ─────────────────────────────────────────

    def _chart_monthly(self, ax, labor, expenses):
        ax.set_title("Monthly Spend Breakdown", pad=8)
        ax.yaxis.set_major_formatter(plt.FuncFormatter(_money_formatter))

        monthly: dict[str, pd.Series] = {}

        if "labor_cost" in labor.columns and "date" in labor.columns:
            tmp = labor.copy()
            tmp["month"] = tmp["date"].dt.to_period("M")
            monthly["Labor"] = tmp.groupby("month")["labor_cost"].sum()

        if "cost" in expenses.columns and "date" in expenses.columns:
            if "category" in expenses.columns:
                for cat in expenses["category"].dropna().unique():
                    tmp = expenses[expenses["category"] == cat].copy()
                    tmp["month"] = tmp["date"].dt.to_period("M")
                    monthly[str(cat)] = tmp.groupby("month")["cost"].sum()
            else:
                tmp = expenses.copy()
                tmp["month"] = tmp["date"].dt.to_period("M")
                monthly["Expenses"] = tmp.groupby("month")["cost"].sum()

        if not monthly:
            ax.text(0.5, 0.5, "No monthly data",
                    transform=ax.transAxes, ha="center", va="center",
                    color=C["dim"], fontsize=10)
            return

        all_months = sorted({m for s in monthly.values() for m in s.index})
        if not all_months:
            return

        x = np.arange(len(all_months))
        month_labels = [str(m) for m in all_months]
        bottom = np.zeros(len(all_months))

        for i, (name, series) in enumerate(monthly.items()):
            vals = np.array([float(series.get(m, 0)) for m in all_months])
            color = C["cat"].get(name, FALLBACK_COLORS[i % len(FALLBACK_COLORS)])
            ax.bar(x, vals, bottom=bottom, label=name, color=color, alpha=0.88)
            bottom += vals

        ax.set_xticks(x)
        ax.set_xticklabels(month_labels, rotation=30, ha="right", fontsize=7)
        ax.set_ylabel("Spend", fontsize=8)
        ax.legend(fontsize=7, facecolor=C["bg"], labelcolor=C["text"],
                  edgecolor=C["border"])

    # ── template generator ────────────────────────────────────────────────────

    def _create_template(self):
        path = filedialog.asksaveasfilename(
            title="Save Template",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="financial_data_template.xlsx")
        if not path:
            return

        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment

            wb = openpyxl.Workbook()

            sheets_data = [
                (
                    "Labor",
                    ["Person", "Project", "Date", "Hours", "Hourly Rate ($)"],
                    [
                        ["Alice Smith",  "Project Alpha", "2025-01-06", 8, 75],
                        ["Bob Jones",    "Project Alpha", "2025-01-06", 6, 85],
                        ["Alice Smith",  "Project Beta",  "2025-01-07", 4, 75],
                        ["Carol White",  "Project Beta",  "2025-01-07", 8, 90],
                        ["Bob Jones",    "Project Gamma", "2025-01-08", 7, 85],
                    ],
                ),
                (
                    "Expenses",
                    ["Project", "Cost", "Date", "Category"],
                    [
                        ["Project Alpha", 1500.00, "2025-01-10", "Material Costs"],
                        ["Project Alpha",  800.00, "2025-01-12", "Purchased Services"],
                        ["Project Beta",   250.00, "2025-01-14", "T&L"],
                        ["Project Beta",  3200.00, "2025-01-15", "Material Costs"],
                        ["Project Gamma",  600.00, "2025-01-09", "Purchased Services"],
                        ["Project Gamma",  150.00, "2025-01-11", "T&L"],
                    ],
                ),
                (
                    "Budgets",
                    ["Project", "Max Budget ($)"],
                    [
                        ["Project Alpha", 50_000],
                        ["Project Beta",  30_000],
                        ["Project Gamma", 20_000],
                    ],
                ),
            ]

            header_fill = PatternFill(
                start_color="7C3AED", end_color="7C3AED", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")

            first = True
            for title, headers, rows in sheets_data:
                ws = wb.active if first else wb.create_sheet(title)
                if first:
                    ws.title = title
                    first = False

                for ci, h in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=ci, value=h)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")

                for ri, row in enumerate(rows, 2):
                    for ci, val in enumerate(row, 1):
                        ws.cell(row=ri, column=ci, value=val)

                for col_cells in ws.columns:
                    width = max(len(str(c.value or "")) for c in col_cells) + 4
                    ws.column_dimensions[col_cells[0].column_letter].width = width

            wb.save(path)
            messagebox.showinfo("Template Saved",
                                f"Template written to:\n{path}\n\n"
                                "Fill in your data and load it with 'Load Excel File'.")
        except ImportError:
            messagebox.showerror(
                "Missing Dependency",
                "openpyxl is required to create the template.\n"
                "Run:  pip install openpyxl")
        except Exception as exc:
            messagebox.showerror("Error", str(exc))


# ── entry point ───────────────────────────────────────────────────────────────

def main():
    root = tk.Tk()

    style = ttk.Style(root)
    try:
        style.theme_use("clam")
    except Exception:
        pass
    style.configure("TCombobox",
                    fieldbackground=C["panel"],
                    background=C["panel"],
                    foreground=C["text"],
                    selectbackground=C["accent"],
                    selectforeground=C["text"])

    app = FinancialDashboard(root)
    root.mainloop()


if __name__ == "__main__":
    main()
