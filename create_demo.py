"""
Generates a realistic demo Excel file for the Financial Data Dashboard.
Run once to create sample_data.xlsx, then load it in the dashboard.
"""

import random
from datetime import date, timedelta

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

random.seed(42)

# ── configuration ──────────────────────────────────────────────────────────────

PROJECTS = {
    "Project Alpha":  {"budget": 85_000,  "start": date(2025, 1, 6)},
    "Project Beta":   {"budget": 52_000,  "start": date(2025, 2, 3)},
    "Project Gamma":  {"budget": 38_000,  "start": date(2025, 1, 20)},
    "Project Delta":  {"budget": 120_000, "start": date(2025, 3, 10)},
    "Project Epsilon":{"budget": 29_000,  "start": date(2025, 2, 17)},
}

PEOPLE = {
    "Alice Smith":   75,
    "Bob Jones":     85,
    "Carol White":   90,
    "David Brown":   70,
    "Emma Davis":    95,
    "Frank Miller":  80,
}

EXPENSE_CATEGORIES = ["Material Costs", "Purchased Services", "T&L"]

EXPENSE_RANGES = {
    "Material Costs":     (500,  8_000),
    "Purchased Services": (200,  3_500),
    "T&L":                (80,     900),
}

# ── generate labor rows ────────────────────────────────────────────────────────

labor_rows = []
for proj, info in PROJECTS.items():
    proj_people = random.sample(list(PEOPLE.keys()), k=random.randint(3, 5))
    start = info["start"]
    for week in range(12):
        week_start = start + timedelta(weeks=week)
        for day_offset in range(5):          # Mon–Fri
            day = week_start + timedelta(days=day_offset)
            if day > date(2025, 6, 30):
                break
            for person in proj_people:
                if random.random() < 0.70:   # 70 % chance they log that day
                    hours = round(random.uniform(2, 9), 1)
                    rate  = PEOPLE[person]
                    labor_rows.append([person, proj, day, hours, rate])

# ── generate expense rows ─────────────────────────────────────────────────────

expense_rows = []
for proj, info in PROJECTS.items():
    start = info["start"]
    for week in range(12):
        week_start = start + timedelta(weeks=week)
        exp_date = week_start + timedelta(days=random.randint(0, 4))
        if exp_date > date(2025, 6, 30):
            break
        num_expenses = random.randint(1, 3)
        for _ in range(num_expenses):
            cat   = random.choice(EXPENSE_CATEGORIES)
            lo, hi = EXPENSE_RANGES[cat]
            cost  = round(random.uniform(lo, hi), 2)
            expense_rows.append([proj, cost, exp_date, cat])

# ── build workbook ─────────────────────────────────────────────────────────────

wb = openpyxl.Workbook()

HEADER_FILL   = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
ALT_FILL      = PatternFill(start_color="F3F0FF", end_color="F3F0FF", fill_type="solid")
HEADER_FONT   = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
BODY_FONT     = Font(name="Calibri", size=10)
CENTER        = Alignment(horizontal="center", vertical="center")
thin          = Side(style="thin", color="D0D0D0")
BORDER        = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_sheet(ws, title, headers, rows):
    ws.title = title
    ws.row_dimensions[1].height = 22

    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.fill   = HEADER_FILL
        c.font   = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER

    for ri, row in enumerate(rows, 2):
        fill = ALT_FILL if ri % 2 == 0 else None
        for ci, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font   = BODY_FONT
            c.border = BORDER
            c.alignment = CENTER
            if fill:
                c.fill = fill

    # auto-width
    for col_cells in ws.columns:
        vals = [str(c.value or "") for c in col_cells]
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(len(v) for v in vals) + 4, 30)

    # freeze header
    ws.freeze_panes = "A2"


# Labor sheet
ws_labor = wb.active
write_sheet(ws_labor, "Labor",
            ["Person", "Project", "Date", "Hours", "Hourly Rate ($)"],
            labor_rows)

# Expenses sheet
ws_exp = wb.create_sheet()
write_sheet(ws_exp, "Expenses",
            ["Project", "Cost", "Date", "Category"],
            expense_rows)

# Budgets sheet
ws_bud = wb.create_sheet()
budget_rows = [[proj, info["budget"]] for proj, info in PROJECTS.items()]
write_sheet(ws_bud, "Budgets",
            ["Project", "Max Budget ($)"],
            budget_rows)

out = "sample_data.xlsx"
wb.save(out)

# ── summary ───────────────────────────────────────────────────────────────────

import pandas as pd
df_labor = pd.DataFrame(labor_rows, columns=["Person","Project","Date","Hours","Rate"])
df_exp   = pd.DataFrame(expense_rows, columns=["Project","Cost","Date","Category"])

df_labor["LaborCost"] = df_labor["Hours"] * df_labor["Rate"]
total_labor   = df_labor["LaborCost"].sum()
total_expense = df_exp["Cost"].sum()

print(f"Demo file written: {out}")
print(f"  Labor rows    : {len(labor_rows)}")
print(f"  Expense rows  : {len(expense_rows)}")
print(f"  Projects      : {len(PROJECTS)}")
print(f"  People        : {len(PEOPLE)}")
print(f"  Total labour  : ${total_labor:,.2f}")
print(f"  Total expenses: ${total_expense:,.2f}")
print(f"  Grand total   : ${total_labor + total_expense:,.2f}")
